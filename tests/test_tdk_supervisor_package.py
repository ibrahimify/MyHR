import re
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.generate_tdk_supervisor_package import (
    CURRENT_ARTIFACT_DIR,
    CURRENT_ARTIFACT_ID,
    TITLE,
    add_inline_markdown,
    append_generated_verification_record,
    build_decision_sheet,
    finalize_visual_inspection,
    write_figure_data,
    write_package_manifest,
)


REPO_ROOT = Path(__file__).resolve().parents[1]
PAPER_SOURCE = REPO_ROOT / "docs" / "TDK_PAPER_DRAFT.md"
PACKAGE_SCRIPT = REPO_ROOT / "scripts" / "generate_tdk_supervisor_package.py"


class TdkSupervisorPackageTests(unittest.TestCase):
    def test_manuscript_source_has_stable_title_and_complete_figures(self):
        text = PAPER_SOURCE.read_text(encoding="utf-8")

        self.assertEqual(
            TITLE,
            "MyHR: Policy-Aware Machine Learning for Anomaly Detection in Workforce Progression",
        )
        self.assertTrue(text.startswith(f"# {TITLE}\n"))
        figures = re.findall(r"\[\[FIGURE:([^|]+)\|([^\]]+)\]\]", text)
        self.assertEqual(11, len(figures))
        self.assertEqual(len(figures), len({name for name, _ in figures}))
        self.assertIn(("primary_recall.png", "Figure 8. Primary exploratory recall comparison in repaired artifact 4feea4508de1bb05. The interval crosses zero; the result does not demonstrate superiority or equivalence."), figures)
        self.assertNotIn("independent_ap.png", text)
        self.assertNotIn("**Keywords:**", text)
        for marker in ("TODO", "TBD", "PLACEHOLDER"):
            self.assertNotIn(marker, text)

    def test_manuscript_preserves_scientific_and_holdout_boundaries(self):
        text = PAPER_SOURCE.read_text(encoding="utf-8")

        required_statements = (
            "No policy-aware advantage was demonstrated",
            CURRENT_ARTIFACT_ID,
            "does not establish equivalence",
            "The final holdout has not been generated or accessed.",
        )
        for statement in required_statements:
            self.assertIn(statement, text)
        self.assertIn("0 of 72", text)
        self.assertIn("Only that `8/72` false-positive interpretation is invalidated", text)
        self.assertIn("designed specification counterexample", text)
        self.assertIn("`deterministic_rules` result is `0/72`", text)
        self.assertIn("does not preserve employee-specific entitlement evidence", text)
        self.assertNotIn("Both interpretations are invalidated", text)
        self.assertIn("explicit evaluation universe", text)
        self.assertIn("this is a design inference applied to the present experiment", text)
        self.assertNotRegex(text, r"(?i)policy-aware (?:isolation forest )?(?:outperforms|is superior)")
        self.assertNotRegex(text, r"(?i)run[_ -]purpose\s*[=:]\s*holdout")
        package_script = PACKAGE_SCRIPT.read_text(encoding="utf-8")
        self.assertIn("39/72 is a constructed specification counterexample", package_script)
        self.assertIn("8/72 was invalidated", package_script)
        self.assertEqual(
            "independent_development_4feea4508de1bb05",
            CURRENT_ARTIFACT_DIR,
        )

    def test_decision_sheet_is_one_landscape_page_source_with_pending_decisions(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "decision.docx"
            build_decision_sheet(output)
            document = Document(output)

        section = document.sections[0]
        self.assertGreater(section.page_width, section.page_height)
        self.assertEqual(1, len(document.tables))
        self.assertEqual(8, len(document.tables[0].rows))
        all_text = "\n".join(
            cell.text for row in document.tables[0].rows for cell in row.cells
        )
        self.assertIn("Scenario proposals", all_text)
        self.assertIn("Freeze and holdout", all_text)
        self.assertIn("[ ] Approve", all_text)

    def test_inline_markdown_is_rendered_without_literal_source_markers(self):
        document = Document()
        paragraph = document.add_paragraph()
        add_inline_markdown(paragraph, "Artifact `abc123` is **verified**.")

        self.assertEqual("Artifact abc123 is verified.", paragraph.text)
        self.assertEqual("Consolas", paragraph.runs[1].font.name)
        self.assertTrue(paragraph.runs[3].bold)

    def test_package_manifest_and_figure_data_are_self_contained(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            (output / "evidence.txt").write_text("verified\n", encoding="utf-8")
            write_figure_data(output)
            write_package_manifest(output, "abc123")
            manifest = json.loads((output / "PACKAGE_MANIFEST.json").read_text(encoding="utf-8"))

        self.assertEqual("abc123", manifest["package_source_revision"])
        self.assertIn("evidence.txt", manifest["files"])
        self.assertIn("figure_data.csv", manifest["files"])
        self.assertNotIn("PACKAGE_MANIFEST.json", manifest["files"])

    def test_generated_verification_record_requires_post_render_finalization(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            report = output / "Verification_Report.md"
            report.write_text("# Verification\n", encoding="utf-8")
            for name in (
                "MyHR_TDK_Paper_Draft_EN.docx",
                "MyHR_TDK_Paper_Draft_EN.pdf",
                "Supervisor_Decision_Sheet.pdf",
            ):
                (output / name).write_bytes(name.encode("ascii"))
            page_report = {
                "total_pages": 39,
                "intro_page": 3,
                "references_page": 37,
                "main_text_pages": 34,
            }
            append_generated_verification_record(report, output, "abc123", page_report)
            write_package_manifest(output, "abc123")
            pending = report.read_text(encoding="utf-8")
            self.assertIn("Visual inspection: PENDING", pending)
            self.assertIn("Main text: pages 3-36, 34 pages", pending)

            finalize_visual_inspection(output, "Every page was inspected.")
            finalized = report.read_text(encoding="utf-8")
            manifest = json.loads((output / "PACKAGE_MANIFEST.json").read_text(encoding="utf-8"))

        self.assertIn("Visual inspection: CONFIRMED. Every page was inspected.", finalized)
        self.assertNotIn("Visual inspection: PENDING", finalized)
        self.assertEqual("abc123", manifest["package_source_revision"])
        self.assertIn("Verification_Report.md", manifest["files"])


if __name__ == "__main__":
    unittest.main()
