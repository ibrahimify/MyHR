from __future__ import annotations

import argparse
import csv
import json
import math
import re
import shutil
import subprocess
import tempfile
import textwrap
from hashlib import sha256
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


TITLE = "MyHR: Policy-Aware Machine Learning for Anomaly Detection in Workforce Progression"
AUTHOR = "Muhammad Ibrahim Shoeb"
SUPERVISOR = "Dr. Husam Al-Magsoosi"
GREEN = "0B6B45"
LIGHT_GREEN = "E8F5EE"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
CURRENT_ARTIFACT_ID = "4feea4508de1bb05"
CURRENT_ARTIFACT_DIR = f"independent_development_{CURRENT_ARTIFACT_ID}"
VALIDATION_COUNTS = {
    "Rules": {"true_positives": 351, "false_positives": 0, "false_negatives": 9},
    "Rules + raw IF": {"true_positives": 357, "false_positives": 51, "false_negatives": 3},
    "Rules + policy IF": {"true_positives": 358, "false_positives": 50, "false_negatives": 2},
}
POLICY_ALERT_HISTORY = {
    "Historical": {"alerts": 39, "denominator": 72},
    "Intermediate": {"alerts": 8, "denominator": 72},
    "Repaired": {"alerts": 0, "denominator": 72},
}


def _save_figure(path: Path) -> None:
    plt.tight_layout()
    plt.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    if path.suffix.lower() == ".png":
        plt.savefig(path.with_suffix(".svg"), bbox_inches="tight", facecolor="white")
    plt.close()


def _diagram_box(ax, x, y, width, height, text, color=LIGHT_GREEN, edge=GREEN):
    rect = plt.Rectangle((x, y), width, height, facecolor=f"#{color}", edgecolor=f"#{edge}", linewidth=1.5)
    ax.add_patch(rect)
    ax.text(x + width / 2, y + height / 2, text, ha="center", va="center", fontsize=10, wrap=True)


def generate_figures(repo: Path, assets: Path) -> None:
    assets.mkdir(parents=True, exist_ok=True)

    fig, ax = plt.subplots(figsize=(10, 4.8))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5)
    ax.axis("off")
    boxes = [
        (0.2, 3.2, 1.7, 1.0, "Administrative\nrecords"),
        (2.3, 3.2, 1.7, 1.0, "Deterministic\nvalidation"),
        (4.4, 3.2, 1.7, 1.0, "Raw and policy-aware\nfeatures"),
        (6.5, 3.2, 1.4, 1.0, "Isolation\nForest"),
        (8.3, 3.2, 1.5, 1.0, "Review\npriority"),
        (3.2, 0.8, 2.0, 1.0, "Evidence-rich\nexplanation"),
        (6.2, 0.8, 2.0, 1.0, "Authorized human\nreview"),
    ]
    for box in boxes:
        _diagram_box(ax, *box)
    arrows = [((1.9, 3.7), (2.3, 3.7)), ((4.0, 3.7), (4.4, 3.7)), ((6.1, 3.7), (6.5, 3.7)), ((7.9, 3.7), (8.3, 3.7)), ((3.2, 3.2), (4.1, 1.8)), ((8.9, 3.2), (7.5, 1.8)), ((5.2, 1.3), (6.2, 1.3))]
    for start, end in arrows:
        ax.annotate("", xy=end, xytext=start, arrowprops={"arrowstyle": "->", "color": "#4A4A4A", "lw": 1.3})
    ax.text(5, 4.65, "Rules-first, ML-second decision-support architecture", ha="center", fontsize=13, fontweight="bold")
    _save_figure(assets / "architecture.png")

    screenshot = repo / "docs" / "media" / "screenshots" / "employee.png"
    if screenshot.exists():
        shutil.copy2(screenshot, assets / "myhr_ui.png")
    else:
        fig, ax = plt.subplots(figsize=(10, 5))
        ax.axis("off")
        ax.text(0.5, 0.5, "MyHR employee performance interface", ha="center", va="center", fontsize=18)
        _save_figure(assets / "myhr_ui.png")

    fig, ax = plt.subplots(figsize=(10, 4.2))
    ax.set_xlim(0, 48)
    ax.set_ylim(-1.2, 3.5)
    ax.axis("off")
    ax.hlines(1.0, 0, 48, color="#333333", linewidth=2)
    for month, label in [(0, "Role start"), (12, "Annual\nincrement"), (24, "Commendation\n+3 months"), (33, "Sanction\n+3 months delay"), (36, "Policy\nchange"), (39, "Decision")]:
        ax.vlines(month, 0.75, 1.25, color=f"#{GREEN}", linewidth=2)
        ax.text(month, 0.55 if month % 24 else 1.45, label, ha="center", va="top" if month % 24 else "bottom", fontsize=10)
    ax.annotate("Effective policy must be resolved at the decision date", xy=(39, 1.0), xytext=(20, 3.0), arrowprops={"arrowstyle": "->", "color": "#B42318"}, fontsize=11, color="#B42318")
    ax.set_title("Progression evidence across service and policy time", fontsize=13, fontweight="bold")
    _save_figure(assets / "policy_timeline.png")

    fig, ax = plt.subplots(figsize=(10, 4.2))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 4)
    ax.axis("off")
    stages = [(0.2, "Development\ninspectable"), (2.25, "Validation\nthresholds and design"), (4.65, "Human protocol\napproval"), (6.75, "Frozen\nrevision and analysis"), (8.45, "One-time\nholdout")]
    for x, label in stages:
        color = "FDECEC" if "holdout" in label.lower() else LIGHT_GREEN
        edge = "B42318" if "holdout" in label.lower() else GREEN
        _diagram_box(ax, x, 1.45, 1.35, 1.0, label, color, edge)
    for x1, x2 in zip([1.55, 3.6, 6.0, 8.1], [2.25, 4.65, 6.75, 8.45]):
        ax.annotate("", xy=(x2, 1.95), xytext=(x1, 1.95), arrowprops={"arrowstyle": "->", "lw": 1.4})
    ax.text(9.1, 0.8, "LOCKED", ha="center", color="#B42318", fontsize=11, fontweight="bold")
    ax.set_title("Research evaluation lifecycle", fontsize=13, fontweight="bold")
    _save_figure(assets / "experiment_flow.png")

    fig, ax = plt.subplots(figsize=(10, 4.8))
    ax.axis("off")
    raw = ["event counts", "raw intervals", "administrative density", "commendation count", "sanction count"]
    policy = ["common core", "service deficit", "deficit ratio", "short-promotion evidence", "policy identity"]
    controls = ["wrong policy", "shuffled policy", "single-feature ablations", "raw without shortcut"]
    for x, title, values, color in [(0.05, "Raw representation", raw, "EAF2FF"), (0.37, "Policy-aware representation", policy, LIGHT_GREEN), (0.72, "Controls", controls, "FFF4D6")]:
        rect = plt.Rectangle((x, 0.12), 0.25, 0.76, transform=ax.transAxes, facecolor=f"#{color}", edgecolor="#555555")
        ax.add_patch(rect)
        ax.text(x + 0.125, 0.81, title, transform=ax.transAxes, ha="center", va="top", fontweight="bold", fontsize=11)
        for i, value in enumerate(values):
            ax.text(x + 0.03, 0.68 - i * 0.12, f"• {value}", transform=ax.transAxes, fontsize=10)
    ax.text(0.5, 0.97, "Matched representations and falsification controls", transform=ax.transAxes, ha="center", fontsize=13, fontweight="bold")
    _save_figure(assets / "feature_design.png")

    fig, ax = plt.subplots(figsize=(10, 3.8))
    ax.axis("off")
    stages = ["Run in owned\nstaging directory", "Write configuration,\npredictions, metrics", "Create SHA-256\nmanifest", "Atomic publish to\nwrite-once directory", "Independent\nrecomputation"]
    for i, label in enumerate(stages):
        x = 0.03 + i * 0.195
        _diagram_box(ax, x, 0.38, 0.16, 0.28, label)
        if i < len(stages) - 1:
            ax.annotate("", xy=(x + 0.19, 0.52), xytext=(x + 0.16, 0.52), xycoords=ax.transAxes, textcoords=ax.transAxes, arrowprops={"arrowstyle": "->", "lw": 1.2})
    ax.text(0.5, 0.85, "Reproducible artifact lifecycle", transform=ax.transAxes, ha="center", fontsize=13, fontweight="bold")
    _save_figure(assets / "artifact_lifecycle.png")

    methods = list(VALIDATION_COUNTS)
    tp = [VALIDATION_COUNTS[method]["true_positives"] for method in methods]
    fp = [VALIDATION_COUNTS[method]["false_positives"] for method in methods]
    fn = [VALIDATION_COUNTS[method]["false_negatives"] for method in methods]
    x = range(len(methods))
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.bar([i - 0.25 for i in x], tp, width=0.25, label="True positives", color="#16834F")
    ax.bar(x, fp, width=0.25, label="False positives", color="#C58A00")
    ax.bar([i + 0.25 for i in x], fn, width=0.25, label="False negatives", color="#C92A2A")
    ax.set_xticks(list(x), methods)
    ax.set_ylabel("Pooled count across 30 validation scenarios")
    ax.legend(frameon=False, ncol=3, loc="upper center")
    ax.grid(axis="y", alpha=0.25)
    _save_figure(assets / "validation_counts.png")

    paired_path = repo / "research_outputs" / "independent_development" / CURRENT_ARTIFACT_DIR / "paired_comparisons.csv"
    paired = pd.read_csv(paired_path)
    result = paired[(paired["method"] == "rule_negative_isolation_forest") & (paired["operating_point"] == "fixed_review_budget")].iloc[0]
    fig, ax = plt.subplots(figsize=(9, 4.8))
    mean = float(result["recall_delta_mean"])
    low = float(result["recall_delta_ci95_low"])
    high = float(result["recall_delta_ci95_high"])
    ax.errorbar([0], [mean], yerr=[[mean - low], [high - mean]], fmt="o", color="#16834F", capsize=8, markersize=9)
    ax.axhline(0, color="#333333", linewidth=1.2)
    ax.set_xlim(-0.8, 0.8)
    ax.set_ylim(-0.06, 0.13)
    ax.set_xticks([0], ["Policy-aware minus raw IF"])
    ax.set_ylabel("Scenario-level recall difference")
    ax.set_title("Primary exploratory comparison at equal review capacity", fontweight="bold")
    ax.text(0, -0.047, "Mean +0.0278; 95% CI [-0.0278, +0.1019]; p=0.6875", ha="center", fontsize=10)
    ax.grid(axis="y", alpha=0.25)
    _save_figure(assets / "primary_recall.png")

    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.axis("off")
    stages = [
        (0.03, "Historical design", "Rule filter consumed\npolicy signal", "FDECEC", "B42318"),
        (0.37, "Audit finding", "Comparison was\nstructurally vacuous", "FFF4D6", "C58A00"),
        (0.71, "Corrected design", "Invariant filter + equal\nresidual ML budget", LIGHT_GREEN, GREEN),
    ]
    for x0, title, body, color, edge in stages:
        _diagram_box(ax, x0, 0.25, 0.26, 0.5, f"{title}\n\n{body}", color, edge)
    for start, end in [((0.29, 0.5), (0.37, 0.5)), ((0.63, 0.5), (0.71, 0.5))]:
        ax.annotate("", xy=end, xytext=start, xycoords=ax.transAxes, textcoords=ax.transAxes, arrowprops={"arrowstyle": "->", "lw": 1.4})
    ax.text(0.5, 0.94, "Design failure preserved as audit evidence", transform=ax.transAxes, ha="center", fontsize=13, fontweight="bold")
    _save_figure(assets / "feature_collapse.png")

    fig, ax = plt.subplots(figsize=(9, 4.8))
    labels = [
        f"{name}\n{item['alerts']}/{item['denominator']}"
        for name, item in POLICY_ALERT_HISTORY.items()
    ]
    values = [item["alerts"] for item in POLICY_ALERT_HISTORY.values()]
    colors = ["#B42318", "#C58A00", "#16834F"]
    bars = ax.bar(labels, values, color=colors)
    ax.set_ylim(0, 45)
    ax.set_ylabel("Deterministic alerts among 72 labelled rare-legitimate cases")
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 1, f"{value}/72", ha="center", fontsize=11, fontweight="bold")
    ax.text(0.5, -0.22, "39/72 is a constructed specification counterexample; 8/72 was invalidated; 0/72 is exploratory evidence.", transform=ax.transAxes, ha="center", fontsize=10)
    ax.grid(axis="y", alpha=0.25)
    _save_figure(assets / "policy_alerts.png")

    fig, ax = plt.subplots(figsize=(10, 4.8))
    ax.axis("off")
    duplicate_stages = ["Partition-local\nrecord set", "Declared pair\nuniverse", "Blocking and\nsimilarity", "Predicted pairs", "TP/FP/FN/TN\nfrom set membership"]
    for i, label in enumerate(duplicate_stages):
        x0 = 0.02 + i * 0.195
        _diagram_box(ax, x0, 0.35, 0.16, 0.3, label)
        if i < len(duplicate_stages) - 1:
            ax.annotate("", xy=(x0 + 0.19, 0.5), xytext=(x0 + 0.16, 0.5), xycoords=ax.transAxes, textcoords=ax.transAxes, arrowprops={"arrowstyle": "->", "lw": 1.2})
    ax.text(0.5, 0.88, "Duplicate-detection evaluation contract", transform=ax.transAxes, ha="center", fontsize=13, fontweight="bold")
    ax.text(0.5, 0.13, "Current repair prevents true negatives from being inferred by capping an incorrectly defined universe.", transform=ax.transAxes, ha="center", fontsize=10)
    _save_figure(assets / "duplicate_results.png")


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER)


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_figure_data(output: Path) -> None:
    rows = []
    for method, values in VALIDATION_COUNTS.items():
        for metric, value in values.items():
            rows.append({
                "figure": "Figure 7",
                "series": method,
                "metric": metric,
                "value": value,
                "denominator": 360,
                "evidence_status": "legacy inherited evidence; fails current provenance schema",
                "source": "legacy/validation_c5878ca7b927ccd2/aggregate_metrics.csv",
            })
    rows.extend([
        {
            "figure": "Figure 8",
            "series": "policy-aware minus raw IF",
            "metric": metric,
            "value": value,
            "denominator": 12,
            "evidence_status": "current verified development artifact",
            "source": "paired_comparisons.csv",
        }
        for metric, value in (
            ("recall_delta_mean", 0.0277777777777777),
            ("recall_delta_ci95_low", -0.0277777777777777),
            ("recall_delta_ci95_high", 0.1018518518518518),
            ("sign_flip_p", 0.6875),
        )
    ])
    sources = {
        "Historical": "legacy/independent_development_e342bd4acf39b06b/aggregate_metrics.csv",
        "Intermediate": "legacy/independent_development_db1a15646d9e76ba/aggregate_metrics.csv",
        "Repaired": "aggregate_metrics.csv",
    }
    statuses = {
        "Historical": "legacy counterexample; incomplete entitlement provenance",
        "Intermediate": "invalidated interpretation; contradictory generated histories",
        "Repaired": "current verified exploratory development evidence",
    }
    for series, values in POLICY_ALERT_HISTORY.items():
        rows.append({
            "figure": "Figure 10",
            "series": series,
            "metric": "rare_legitimate_alerts",
            "value": values["alerts"],
            "denominator": values["denominator"],
            "evidence_status": statuses[series],
            "source": sources[series],
        })
    with (output / "figure_data.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)


def write_package_manifest(output: Path, revision: str) -> None:
    files = {
        str(path.relative_to(output)).replace("\\", "/"): _sha256(path)
        for path in sorted(output.rglob("*"))
        if path.is_file() and path.name != "PACKAGE_MANIFEST.json"
    }
    (output / "PACKAGE_MANIFEST.json").write_text(
        json.dumps(
            {
                "algorithm": "sha256",
                "package_source_revision": revision,
                "files": files,
            },
            indent=2,
            sort_keys=True,
        ) + "\n",
        encoding="utf-8",
    )
def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.5)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(3.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_sizes = {1: 16, 2: 14, 3: 12}
    for level, size in heading_sizes.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in [style.name for style in document.styles]:
        code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        code.font.size = Pt(10)
        code.paragraph_format.left_indent = Cm(0.7)
        code.paragraph_format.right_indent = Cm(0.7)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(6)

    for doc_section in document.sections:
        _add_page_number(doc_section.footer.paragraphs[0])


def add_title_pages(document: Document, lines: list[str]) -> None:
    document.add_paragraph("Budapest University of Technology and Economics", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Faculty of Electrical Engineering and Informatics", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(80)
    run = title.add_run(TITLE)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)
    subtitle = document.add_paragraph("TDK Paper Supervisor Review Draft")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    document.add_paragraph()
    for label, value in [("Author", AUTHOR), ("Supervisor", SUPERVISOR), ("Date", "30 October 2026")]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)
    warning = document.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.paragraph_format.space_before = Pt(80)
    run = warning.add_run("Development evidence only. Draft protocol. Final holdout locked.")
    run.bold = True
    run.font.color.rgb = RGBColor(180, 35, 24)
    document.add_page_break()


def _parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    return rows


def add_inline_markdown(paragraph, text: str) -> None:
    """Render the small inline Markdown subset used by the paper source."""
    cursor = 0
    for match in re.finditer(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(10)
        else:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        cant_split = OxmlElement("w:cantSplit")
        table.rows[row_index]._tr.get_or_add_trPr().append(cant_split)
        if row_index == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            table.rows[row_index]._tr.get_or_add_trPr().append(repeat_header)
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            text = row[column_index] if column_index < len(row) else ""
            cell.text = ""
            add_inline_markdown(cell.paragraphs[0], text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_borders(cell)
            if row_index == 0:
                _set_cell_shading(cell, GREEN)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            elif row_index % 2 == 0:
                _set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    if run.font.name != "Consolas":
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(9.5)
    if len(rows) <= 6:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    document.add_paragraph()


def add_figure(document: Document, assets: Path, name: str, caption: str) -> None:
    path = assets / name
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    width = 5.0 if name == "myhr_ui.png" else 6.0
    run.add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = False
    for run in caption_paragraph.runs:
        run.italic = True
        run.font.size = Pt(10)


def build_docx(source: Path, assets: Path, output: Path) -> None:
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    document = Document()
    configure_document(document)
    add_title_pages(document, lines)

    in_code = False
    code_lines: list[str] = []
    paragraph_lines: list[str] = []
    table_lines: list[str] = []
    seen_intro = False
    seen_references = False
    after_heading = False

    def add_hidden_marker(text: str, bookmark_name: str, bookmark_id: int) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 0.1
        run = paragraph.add_run(text)
        run.font.hidden = True
        run.font.size = Pt(1)
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), str(bookmark_id))
        bookmark_start.set(qn("w:name"), bookmark_name)
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), str(bookmark_id))
        paragraph._p.insert(0, bookmark_start)
        paragraph._p.append(bookmark_end)

    def flush_paragraph() -> None:
        nonlocal paragraph_lines, after_heading
        if paragraph_lines:
            content = " ".join(line.strip() for line in paragraph_lines).strip()
            if content:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if after_heading:
                    paragraph.paragraph_format.keep_together = True
                add_inline_markdown(paragraph, content)
                after_heading = False
            paragraph_lines = []

    def flush_table() -> None:
        nonlocal table_lines
        if table_lines:
            add_table(document, _parse_table(table_lines))
            table_lines = []

    for line in lines:
        if line.startswith("# "):
            continue
        if line.startswith("**Author:") or line.startswith("**Supervisor:") or line.startswith("**Institution:") or line.startswith("**Paper language:") or line.startswith("**Status:"):
            continue
        if line.strip().startswith("```"):
            flush_paragraph()
            flush_table()
            if in_code:
                paragraph = document.add_paragraph("\n".join(code_lines), style="Code Block")
                for run in paragraph.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(10)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("|"):
            flush_paragraph()
            table_lines.append(line)
            continue
        flush_table()
        figure = re.fullmatch(r"\[\[FIGURE:([^|]+)\|(.+)\]\]", line.strip())
        if figure:
            flush_paragraph()
            add_figure(document, assets, figure.group(1), figure.group(2))
            continue
        if line.startswith("## "):
            flush_paragraph()
            heading = line[3:].strip()
            if heading.startswith("1 Introduction") and not seen_intro:
                document.add_page_break()
                add_hidden_marker("MYHR_MAIN_TEXT_START_2026", "MYHR_MAIN_TEXT_START", 100)
                seen_intro = True
            elif heading == "References" and not seen_references:
                document.add_page_break()
                add_hidden_marker("MYHR_MAIN_TEXT_END_2026", "MYHR_MAIN_TEXT_END", 101)
                seen_references = True
            elif re.match(r"^[278] ", heading):
                document.add_page_break()
            heading_paragraph = document.add_heading(heading, level=1)
            heading_paragraph.paragraph_format.keep_with_next = True
            after_heading = True
            continue
        if line.startswith("### "):
            flush_paragraph()
            heading_paragraph = document.add_heading(line[4:].strip(), level=2)
            heading_paragraph.paragraph_format.keep_with_next = True
            after_heading = True
            continue
        if line.startswith("#### "):
            flush_paragraph()
            heading_paragraph = document.add_heading(line[5:].strip(), level=3)
            heading_paragraph.paragraph_format.keep_with_next = True
            after_heading = True
            continue
        if re.match(r"^\d+\. ", line):
            flush_paragraph()
            paragraph = document.add_paragraph()
            add_inline_markdown(paragraph, line.strip())
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.first_line_indent = Cm(-0.5)
            continue
        if line.startswith("- "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, line[2:].strip())
            continue
        if not line.strip():
            flush_paragraph()
            continue
        paragraph_lines.append(line)

    flush_paragraph()
    flush_table()
    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = TITLE
    document.core_properties.author = AUTHOR
    document.core_properties.subject = "BME VIK TDK 2026 supervisor review draft"
    document.core_properties.keywords = "machine learning, anomaly detection, Isolation Forest, workforce progression, MyHR"
    document.save(output)


def build_decision_sheet(output: Path) -> None:
    document = Document()
    configure_document(document)
    section = document.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MyHR TDK Supervisor Decision Sheet")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    status = document.add_paragraph("Nothing is approved until a human records a decision. The final holdout remains locked.")
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.runs[0].bold = True
    status.runs[0].font.color.rgb = RGBColor(180, 35, 24)
    identity = document.add_paragraph("Date: ____________________    Reviewer: __________________________________________")
    identity.alignment = WD_ALIGN_PARAGRAPH.LEFT
    identity.paragraph_format.space_after = Pt(6)

    rows = [
        ("1. Scenario proposals", "Approve, reject, or modify Proposals 1 and 3-6: lawful late corrections, transitions, eligibility, and matrix composition.", "Recommendation: approve the repaired semantics only after checking the concrete examples."),
        ("2. Grandfathering", "Choose whether entitlement follows employee, event, policy epoch, or an explicit evidence record.", "Recommendation: require explicit entitlement evidence; do not infer it from labels."),
        ("3. Temporal views", "Define knowledge-time and review-time views for late facts and corrections.", "Recommendation: freeze both views before generating protocol-scale evidence."),
        ("4. Expert review", "Name two blinded reviewers and fix sample, questions, agreement statistic, and adjudication rule.", "Recommendation: HR-policy plus information-systems expertise."),
        ("5. Experiment scale", "Approve repeated scenario seeds, population cells, and zero-positive handling for development and validation.", "Recommendation: scenario seeds remain the statistical units."),
        ("6. Multiplicity", "Confirm the secondary comparison family and Holm adjustment before the next evaluation.", "Recommendation: keep recall primary and AP secondary."),
        ("7. Freeze and holdout", "After all above items and artifacts pass audit, grant or withhold a dated protocol freeze and one-time holdout authorization.", "Recommendation: withhold authorization until every prerequisite is signed."),
    ]
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(21.5)
    table.cell(0, 0).text = "Decision"
    table.cell(0, 1).text = "Evidence-based choice"
    for cell in table.rows[0].cells:
        _set_cell_shading(cell, GREEN)
        _set_cell_borders(cell)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for index, (name, choice, recommendation) in enumerate(rows, start=1):
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = f"{choice}\n{recommendation}\n[ ] Approve   [ ] Reject   [ ] Modify: ______________________________"
        for cell in cells:
            _set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2 == 0:
                _set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 0.92
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.2)
        cells[0].paragraphs[0].runs[0].bold = True
    for row in table.rows:
        row.cells[0].width = Cm(5.0)
        row.cells[1].width = Cm(21.5)
    footer = document.add_paragraph("Protocol freeze: [ ] No   [ ] Yes, after recorded modifications\nSignature: ______________________________    Date: ____________________")
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.paragraph_format.space_before = Pt(5)
    footer.runs[0].bold = True
    document.save(output)


def convert_with_word(docx_path: Path, pdf_path: Path, report_path: Path) -> dict:
    ps1 = docx_path.parent / "_convert_tdk_document.ps1"
    script = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {{
    $document = $word.Documents.Open('{str(docx_path).replace("'", "''")}')
    $document.Repaginate()
    foreach ($field in $document.Fields) {{ try {{ $field.Update() | Out-Null }} catch {{ }} }}
    $total = $document.ComputeStatistics(2)
    $introPage = if ($document.Bookmarks.Exists('MYHR_MAIN_TEXT_START')) {{ $document.Bookmarks.Item('MYHR_MAIN_TEXT_START').Range.Information(3) }} else {{ 0 }}
    $refsPage = if ($document.Bookmarks.Exists('MYHR_MAIN_TEXT_END')) {{ $document.Bookmarks.Item('MYHR_MAIN_TEXT_END').Range.Information(3) }} else {{ 0 }}
    $mainPages = if ($introPage -gt 0 -and $refsPage -gt $introPage) {{ $refsPage - $introPage }} else {{ 0 }}
    $document.Save()
    $document.ExportAsFixedFormat('{str(pdf_path).replace("'", "''")}', 17)
    $result = [ordered]@{{ total_pages = $total; intro_page = $introPage; references_page = $refsPage; main_text_pages = $mainPages }}
    $result | ConvertTo-Json | Set-Content -Encoding UTF8 '{str(report_path).replace("'", "''")}'
    $document.Close($false)
}} finally {{
    $word.Quit()
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
}}
"""
    ps1.write_text(script, encoding="utf-8-sig")
    try:
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps1)], check=True)
    finally:
        ps1.unlink(missing_ok=True)
    return json.loads(report_path.read_text(encoding="utf-8-sig"))


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate the MyHR TDK supervisor review package.")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[1]
    output = args.output or (repo.parent / "TDK" / "final" / "2026-10-30-supervisor")
    if output.exists() and any(output.iterdir()):
        archive = output.parent / "archive" / "2026-10-30-pre-autonomous-audit-8d7b4e0"
        if not archive.exists():
            archive.parent.mkdir(parents=True, exist_ok=True)
            shutil.copytree(output, archive)
    output.mkdir(parents=True, exist_ok=True)
    docx_path = output / "MyHR_TDK_Paper_Draft_EN.docx"
    pdf_path = output / "MyHR_TDK_Paper_Draft_EN.pdf"
    page_report_path = output / "page_report.json"
    with tempfile.TemporaryDirectory(prefix="myhr_tdk_assets_") as directory:
        assets = Path(directory)
        generate_figures(repo, assets)
        build_docx(repo / "docs" / "TDK_PAPER_DRAFT.md", assets, docx_path)
        for stale_name in ("independent_ap.png", "independent_ap.svg"):
            (output / "figures" / stale_name).unlink(missing_ok=True)
        shutil.copytree(assets, output / "figures", dirs_exist_ok=True)
    page_report = convert_with_word(docx_path, pdf_path, page_report_path)
    decision_docx = output / "_Supervisor_Decision_Sheet.docx"
    decision_pdf = output / "Supervisor_Decision_Sheet.pdf"
    decision_report = output / "_decision_page_report.json"
    build_decision_sheet(decision_docx)
    convert_with_word(decision_docx, decision_pdf, decision_report)
    decision_docx.unlink(missing_ok=True)
    decision_report.unlink(missing_ok=True)
    page_report_path.unlink(missing_ok=True)
    shutil.copy2(repo / "docs" / "TDK_SUPERVISOR_BRIEF_2026-10-30.md", output / "README.md")
    verification = repo / "docs" / "TDK_SUPERVISOR_PACKAGE_VERIFICATION_2026-10-30.md"
    if verification.exists():
        shutil.copy2(verification, output / "Verification_Report.md")
    artifact = repo / "research_outputs" / "independent_development" / CURRENT_ARTIFACT_DIR
    data_dir = output / "data"
    data_dir.mkdir(exist_ok=True)
    artifact_manifest = json.loads((artifact / "manifest.json").read_text(encoding="utf-8"))
    for name in artifact_manifest["files"]:
        shutil.copy2(artifact / name, data_dir / name)
    legacy_sources = {
        "validation_c5878ca7b927ccd2": repo / "research_outputs" / "protocol_v2_verified" / "validation_c5878ca7b927ccd2",
        "independent_development_e342bd4acf39b06b": repo / "research_outputs" / "independent_development" / "independent_development_e342bd4acf39b06b",
        "independent_development_db1a15646d9e76ba": repo / "research_outputs" / "independent_development" / "independent_development_db1a15646d9e76ba",
    }
    for label, source in legacy_sources.items():
        destination = data_dir / "legacy" / label
        destination.mkdir(parents=True, exist_ok=True)
        for name in ("README.md", "metadata.json", "manifest.json", "aggregate_metrics.csv", "paired_comparisons.csv"):
            shutil.copy2(source / name, destination / name)
    write_figure_data(data_dir)
    revision = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=repo,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    (output / "START_HERE.md").write_text(
        "\n".join([
            "# MyHR TDK Supervisor Review Draft",
            "",
            f"Reviewed branch: `research/policy-aware-anomaly` at `{revision}`.",
            "",
            "## Start Here",
            "",
            "1. Read `README.md` for the research question, verified findings, and limitations.",
            "2. Read `MyHR_TDK_Paper_Draft_EN.pdf` for the complete English manuscript.",
            "3. Use `Supervisor_Decision_Sheet.pdf` for the remaining human decisions.",
            "4. Check `Verification_Report.md` for tests, citations, artifact audit, page count, and hashes.",
            "5. Inspect `figures/` and `data/` for reproducible plots and their underlying tables.",
            "",
            "## Verified Development Evidence",
            "",
            f"- Artifact: `{CURRENT_ARTIFACT_DIR}`.",
            "- 12 scenarios, 864 eligible test histories, 108 anomalies, and 72 rare-legitimate histories.",
            "- Equal 96-record residual queues: raw found 36 anomalies; policy-aware found 39.",
            "- Recall delta: +0.0278, 95% CI [-0.0278, +0.1019], paired sign-flip p=0.6875.",
            "- Interpretation: no advantage demonstrated; equivalence is not established.",
            "",
            "## Remaining Limitations And Decisions",
            "",
            "- Synthetic scenario realism still requires two blinded human reviewers.",
            "- Employee-specific grandfathering is not represented as structured generator evidence.",
            "- Protocol-scale development and validation remain outstanding.",
            "- Protocol freeze, final holdout access, scientific conclusions, PR merge, and final submission require human authorization.",
            "",
            "The final holdout was not generated, inspected, scored, or unlocked.",
            "",
        ]),
        encoding="utf-8",
    )
    write_package_manifest(output, revision)
    print(json.dumps({"docx": str(docx_path), "pdf": str(pdf_path), "decision_sheet": str(decision_pdf), **page_report}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
